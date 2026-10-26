from __future__ import annotations

import csv
import sqlite3
import sys
import webbrowser
from datetime import datetime
from pathlib import Path

import keyring
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from PySide6.QtCore import QThread, Signal, Qt, QSettings, QTimer
from PySide6.QtGui import QAction, QColor, QBrush, QFont
from PySide6.QtWidgets import (
    QApplication, QAbstractItemView, QCheckBox, QComboBox, QDialog, QDialogButtonBox,
    QFileDialog, QFormLayout, QFrame, QGridLayout, QHBoxLayout, QHeaderView, QLabel,
    QLineEdit, QListWidget, QMainWindow, QMessageBox, QPushButton,
    QPlainTextEdit, QScrollArea, QSplitter, QStackedWidget, QTableWidget,
    QTableWidgetItem, QVBoxLayout, QWidget,
)

from db import Database
from github_api import GitHubClient, repository_from_input, request_device_code, poll_device_token

KEYRING_SERVICE = "GitHubVulnTracker"
KEYRING_USER = "github_oauth_token"
KEYRING_REFRESH = "github_oauth_refresh_token"
OAUTH_SCOPE = "repo"
GITHUB_STATES = ["local", "triage", "draft", "published", "closed"]
BOUNTY_STATES = ["local", "submitted", "triage", "accepted", "published", "duplicate", "informative",
                 "not_applicable", "resolved", "closed"]
ALL_STATES = list(dict.fromkeys(GITHUB_STATES + BOUNTY_STATES))
SOURCE_LABELS = {"github": "GitHub (동기화)", "manual": "GitHub (수동)", "bug_bounty": "버그바운티"}
TABLE_SOURCE_LABELS = {"github": "GitHub 연동", "manual": "GitHub 수동", "bug_bounty": "버그바운티"}
REPORT_ID_COLUMN = 10


def fmt_date(value: str | None) -> str:
    if not value:
        return "-"
    try:
        dt = datetime.fromisoformat(value.replace("Z", "+00:00"))
        return dt.strftime("%Y-%m-%d %H:%M")
    except Exception:
        return value


class SyncWorker(QThread):
    progress = Signal(str)
    item = Signal(dict)
    done = Signal(int, int)
    failed = Signal(str)

    def __init__(self, token: str, repos: list[str]):
        super().__init__()
        self.token = token
        self.repos = repos

    def run(self):
        try:
            client = GitHubClient(self.token)
            login = client.current_user()
            total = failures = 0
            for repo in self.repos:
                self.progress.emit(f"동기화 중: {repo}")
                try:
                    for report in client.sync_repository(repo, login):
                        self.item.emit(report)
                        total += 1
                except Exception as e:
                    failures += 1
                    self.progress.emit(f"실패: {repo} ({e})")
            self.done.emit(total, failures)
        except Exception as e:
            self.failed.emit(str(e))


class OAuthWorker(QThread):
    code_ready = Signal(str, str, int)
    authenticated = Signal(str, str, str)
    failed = Signal(str)

    def __init__(self, client_id: str):
        super().__init__()
        self.client_id = client_id
        self._stop = False

    def cancel(self):
        self._stop = True

    def run(self):
        try:
            device = request_device_code(self.client_id, OAUTH_SCOPE)
            self.code_ready.emit(device["verification_uri"], device["user_code"], int(device.get("expires_in", 900)))
            token_data = poll_device_token(
                self.client_id, device["device_code"], int(device.get("interval", 5)),
                int(device.get("expires_in", 900)), should_stop=lambda: self._stop
            )
            token = token_data["access_token"]
            login = GitHubClient(token).current_user()
            self.authenticated.emit(login, token, token_data.get("refresh_token") or "")
        except Exception as e:
            self.failed.emit(str(e))


class ReportDialog(QDialog):
    def __init__(self, parent=None, report=None, source="manual"):
        super().__init__(parent)
        self.report = dict(report) if report is not None else {}
        self.setWindowTitle("취약점 제보 수정" if self.report else "취약점 제보 추가")
        self.setMinimumWidth(560)
        self.form = QFormLayout(self)
        self.source = QComboBox()
        self.source.addItem("GitHub (수동)", "manual")
        self.source.addItem("버그바운티", "bug_bounty")
        self.repo = QLineEdit()
        self.platform = QComboBox(); self.platform.setEditable(True)
        self.platform.addItems(["", "HackerOne", "Bugcrowd", "Intigriti", "YesWeHack", "직접 제보", "기타"])
        self.program = QLineEdit(); self.program.setPlaceholderText("프로그램 또는 기업 이름 (선택)")
        self.report_reference = QLineEdit(); self.report_reference.setPlaceholderText("제보 번호 / Report ID (선택)")
        self.title = QLineEdit(); self.title.setPlaceholderText("제보 제목을 입력하세요")
        self.ghsa = QLineEdit()
        self.cve = QLineEdit()
        self.state = QComboBox()
        self.state_help = QLabel("accepted = 제보 인정  ·  published = 제보 공개 (Published 집계)")
        self.state_help.setWordWrap(True)
        self.manual_help = QLabel("수동 GitHub 제보를 저장하면 해당 저장소가 동기화 목록에도 추가됩니다.")
        self.manual_help.setWordWrap(True)
        self.severity = QComboBox(); self.severity.addItems(["", "none", "low", "moderate", "medium", "high", "critical"])
        self.url = QLineEdit(); self.url.setPlaceholderText("제보 페이지 URL (선택)")
        self.notes = QPlainTextEdit(); self.notes.setMaximumHeight(100)
        for label, widget in [("출처", self.source), ("대상 *", self.repo),
                              ("플랫폼", self.platform), ("프로그램", self.program),
                              ("제보 번호", self.report_reference), ("제목 *", self.title), ("GHSA", self.ghsa),
                              ("CVE", self.cve), ("상태", self.state), ("상태 기준", self.state_help),
                              ("GitHub 안내", self.manual_help), ("Severity", self.severity),
                              ("URL", self.url), ("메모", self.notes)]:
            self.form.addRow(label, widget)
        self.source.setCurrentIndex(self.source.findData(self.report.get("source", source)))
        self._source_changed()
        self.source.currentIndexChanged.connect(self._source_changed)
        for widget, key in [(self.repo, "repository"), (self.program, "program"),
                            (self.report_reference, "report_reference"), (self.title, "title"),
                            (self.ghsa, "ghsa_id"), (self.cve, "cve_id"), (self.url, "html_url")]:
            widget.setText(self.report.get(key) or "")
        self.platform.setCurrentText(self.report.get("platform") or "")
        for combo, key in [(self.state, "state"), (self.severity, "severity")]:
            value = self.report.get(key) or ""
            if value and combo.findText(value) < 0:
                combo.addItem(value)
            if value:
                combo.setCurrentText(value)
        self.notes.setPlainText(self.report.get("notes") or "")
        buttons = QDialogButtonBox(QDialogButtonBox.Save | QDialogButtonBox.Cancel)
        buttons.button(QDialogButtonBox.Save).setText("저장")
        buttons.button(QDialogButtonBox.Cancel).setText("취소")
        buttons.accepted.connect(self.accept); buttons.rejected.connect(self.reject)
        self.form.addRow(buttons)
        self.repo.setFocus()

    def _source_changed(self):
        bounty = self.source.currentData() == "bug_bounty"
        for widget in [self.platform, self.program, self.report_reference]:
            self.form.setRowVisible(widget, bounty)
        self.form.setRowVisible(self.ghsa, not bounty)
        self.form.setRowVisible(self.state_help, bounty)
        self.form.setRowVisible(self.manual_help, not bounty)
        self.form.labelForField(self.repo).setText("대상 *" if bounty else "저장소 *")
        self.repo.setPlaceholderText("서비스 / 도메인 / 앱 이름" if bounty else "owner/repository 또는 GitHub URL")
        previous = self.state.currentText()
        self.state.clear(); self.state.addItems(BOUNTY_STATES if bounty else GITHUB_STATES)
        if not self.report and bounty and previous in {"", "local"}:
            self.state.setCurrentText("submitted")
        elif previous in [self.state.itemText(i) for i in range(self.state.count())]:
            self.state.setCurrentText(previous)

    def accept(self):
        if not self.repo.text().strip():
            QMessageBox.warning(self, "필수값", "대상 또는 저장소를 입력하세요.")
            self.repo.setFocus()
            return
        if not self.title.text().strip():
            QMessageBox.warning(self, "필수값", "제보 제목을 입력하세요.")
            self.title.setFocus()
            return
        if self.source.currentData() == "manual":
            repo = repository_from_input(self.repo.text())
            if not repo:
                QMessageBox.warning(self, "저장소 형식", "owner/repository 또는 GitHub 저장소 URL을 입력하세요.")
                return
            self.repo.setText(repo)
        super().accept()

    def data(self):
        now = datetime.now().isoformat(timespec="seconds")
        bounty = self.source.currentData() == "bug_bounty"
        return {
            "repository": self.repo.text().strip(), "title": self.title.text().strip(),
            "ghsa_id": None if bounty else self.ghsa.text().strip() or None, "cve_id": self.cve.text().strip() or None,
            "state": self.state.currentText(), "severity": self.severity.currentText() or None,
            "html_url": self.url.text().strip() or None, "notes": self.notes.toPlainText(),
            "submitted_at": self.report.get("submitted_at") if self.report else now,
            "updated_at": now, "source": self.source.currentData(),
            "platform": self.platform.currentText().strip() if bounty else "",
            "program": self.program.text().strip() if bounty else "",
            "report_reference": self.report_reference.text().strip() if bounty else "",
        }


class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.db = Database()
        self.worker = None
        self.current_report_id = None
        self._closing_pending = False
        self._loading_detail = False
        self._detail_dirty = False
        self._detail_save_timer = QTimer(self)
        self._detail_save_timer.setSingleShot(True)
        self._detail_save_timer.setInterval(800)
        self._detail_save_timer.timeout.connect(self._autosave_detail)
        self.settings = QSettings("LocalTools", "GitHubVulnTracker")
        self.auth_worker = None
        self.setWindowTitle("Vulnerability Tracker v0.4")
        self.resize(1220, 700)
        self.setMinimumSize(980, 620)
        self._build_ui(); self._apply_style(); self.refresh_all(); self._update_report_layout()
        self.statusBar().showMessage("준비")

    def _build_ui(self):
        root = QWidget(); self.setCentralWidget(root)
        outer = QHBoxLayout(root); outer.setContentsMargins(0, 0, 0, 0); outer.setSpacing(0)
        nav = QFrame(); nav.setObjectName("nav"); nav.setFixedWidth(180)
        nl = QVBoxLayout(nav)
        brand = QLabel("VULN\nTRACKER"); brand.setObjectName("brand"); nl.addWidget(brand)
        self.btn_dashboard = QPushButton("대시보드")
        self.btn_reports = QPushButton("제보 목록")
        self.btn_settings = QPushButton("GitHub 설정")
        for i, b in enumerate([self.btn_dashboard, self.btn_reports, self.btn_settings]):
            b.setCheckable(True); b.clicked.connect(lambda _, idx=i: self.show_page(idx)); nl.addWidget(b)
        nl.addStretch(1); nl.addWidget(QLabel("GitHub · Bug Bounty · v0.4")); outer.addWidget(nav)
        self.stack = QStackedWidget(); self.stack.addWidget(self._dashboard_page()); self.stack.addWidget(self._reports_page()); self.stack.addWidget(self._settings_page())
        outer.addWidget(self.stack, 1); self.show_page(0)
        a = QAction("취약점 추가", self); a.setShortcut("Ctrl+N"); a.triggered.connect(self.add_report); self.addAction(a)
        e = QAction("Excel 내보내기", self); e.setShortcut("Ctrl+E"); e.triggered.connect(self.export_xlsx); self.addAction(e)
        find = QAction("제보 검색", self); find.setShortcut("Ctrl+F"); find.triggered.connect(self.focus_search); self.addAction(find)
        save = QAction("메모 저장", self); save.setShortcut("Ctrl+S"); save.triggered.connect(self.save_detail); self.addAction(save)

    def _dashboard_page(self):
        w = QWidget(); l = QVBoxLayout(w); l.setContentsMargins(28,24,28,24)
        title = QLabel("제보 현황"); title.setObjectName("pageTitle"); l.addWidget(title)
        cards = QGridLayout()
        self.card_total = self._stat_card(cards, "전체", 0, 0, lambda: self.show_reports_for())
        self.card_bounty = self._stat_card(cards, "버그바운티", 0, 1, lambda: self.show_reports_for(source="bug_bounty"))
        self.card_triage = self._stat_card(cards, "검토 중 · triage", 0, 2, lambda: self.show_reports_for(state="triage"))
        self.card_draft = self._stat_card(cards, "초안 · draft", 1, 0, lambda: self.show_reports_for(state="draft"))
        self.card_published = self._stat_card(cards, "공개 · published", 1, 1, lambda: self.show_reports_for(state="published"))
        self.card_cve = self._stat_card(cards, "CVE 발급", 1, 2, lambda: self.show_reports_for(cve=True))
        l.addLayout(cards)
        h = QHBoxLayout()
        for text, fn in [("+ GitHub 제보", self.add_report), ("+ 버그바운티", lambda: self.add_report(source="bug_bounty")), ("GitHub 링크 추가", self.add_repo_from_link_dialog), ("GitHub 동기화", self.sync_github), ("Excel 전체 내보내기", self.export_xlsx)]:
            b = QPushButton(text); b.clicked.connect(fn); h.addWidget(b)
        h.addStretch(); l.addLayout(h)
        lab = QLabel("최근 업데이트"); lab.setObjectName("sectionTitle"); l.addWidget(lab)
        self.recent_table = self._make_table()
        self.recent_table.cellDoubleClicked.connect(self.open_recent_report)
        self.recent_table.setToolTip("항목을 두 번 클릭하면 제보 상세 화면으로 이동합니다.")
        l.addWidget(self.recent_table)
        return w

    def _stat_card(self, layout, label, row, column, action):
        card = QPushButton(); card.setObjectName("statCard"); card.setAccessibleName(f"{label} 제보 보기")
        card.setCursor(Qt.PointingHandCursor)
        card.setMinimumHeight(78)
        inner = QVBoxLayout(card); inner.setContentsMargins(14, 8, 14, 8)
        value = QLabel("0"); value.setObjectName("statValue")
        value.setAttribute(Qt.WA_TransparentForMouseEvents)
        caption = QLabel(label); caption.setAttribute(Qt.WA_TransparentForMouseEvents)
        inner.addWidget(value); inner.addWidget(caption)
        card.clicked.connect(action); layout.addWidget(card, row, column)
        return value

    def _make_table(self):
        t = QTableWidget(0, 11)
        t.setHorizontalHeaderLabels(["출처", "대상 / 저장소", "제목", "상태", "Severity", "CVE", "제보 번호", "GHSA", "수정", "변경", "ID"])
        for column in (4, 5, 6, 7, 9, REPORT_ID_COLUMN):
            t.setColumnHidden(column, True)
        t.setSelectionBehavior(QAbstractItemView.SelectRows); t.setSelectionMode(QAbstractItemView.SingleSelection); t.setEditTriggers(QAbstractItemView.NoEditTriggers); t.verticalHeader().setVisible(False)
        t.setAlternatingRowColors(True)
        t.horizontalHeader().setSectionResizeMode(QHeaderView.Interactive)
        for column, width in [(0, 95), (1, 125), (3, 85), (8, 110)]:
            t.setColumnWidth(column, width)
        t.horizontalHeader().setSectionResizeMode(2, QHeaderView.Stretch)
        return t

    def _reports_page(self):
        w = QWidget(); l = QVBoxLayout(w); l.setContentsMargins(28,24,28,24)
        title = QLabel("제보 목록"); title.setObjectName("pageTitle"); l.addWidget(title)
        search_row = QHBoxLayout()
        self.search = QLineEdit(); self.search.setPlaceholderText("대상 / 플랫폼 / 프로그램 / 제보 번호 / 제목 / GHSA / CVE 검색"); self.search.setClearButtonEnabled(True); self.search.textChanged.connect(self.refresh_reports)
        search_row.addWidget(QLabel("검색")); search_row.addWidget(self.search, 1)
        reset_button = QPushButton("필터 초기화"); reset_button.clicked.connect(self.clear_filters); search_row.addWidget(reset_button)
        l.addLayout(search_row)
        controls = QHBoxLayout()
        self.state_filter = QComboBox(); self.state_filter.addItems(["전체"] + ALL_STATES); self.state_filter.currentTextChanged.connect(self.refresh_reports)
        self.source_filter = QComboBox(); self.source_filter.addItem("모든 출처", ""); self.source_filter.addItem("GitHub", "github"); self.source_filter.addItem("버그바운티", "bug_bounty"); self.source_filter.currentIndexChanged.connect(self._source_filter_changed)
        self.cve_filter = QCheckBox("CVE 있음"); self.cve_filter.toggled.connect(self.refresh_reports)
        controls.addWidget(QLabel("출처")); controls.addWidget(self.source_filter); controls.addWidget(QLabel("상태")); controls.addWidget(self.state_filter); controls.addWidget(self.cve_filter)
        controls.addStretch(1)
        for text, fn in [("+ GitHub 제보", self.add_report), ("+ 버그바운티", lambda: self.add_report(source="bug_bounty")), ("CSV", self.export_csv), ("Excel", self.export_xlsx)]:
            b = QPushButton(text); b.clicked.connect(fn); controls.addWidget(b)
        l.addLayout(controls)
        self.result_count = QLabel("검색 결과 0건")
        l.addWidget(self.result_count)
        splitter = QSplitter(Qt.Horizontal); splitter.setChildrenCollapsible(False)
        self.report_splitter = splitter
        self.report_table = self._make_table(); self.report_table.itemSelectionChanged.connect(self.show_selected_report); self.report_table.cellDoubleClicked.connect(self.focus_selected_detail); self.report_table.setToolTip("한 번 클릭하면 상세 정보가 보입니다. 두 번 클릭하면 메모로 이동합니다."); splitter.addWidget(self.report_table)
        detail = QFrame(); detail.setObjectName("detail"); detail_layout = QVBoxLayout(detail)
        detail_content = QWidget(); detail_content.setObjectName("detailContent")
        dl = QVBoxLayout(detail_content)
        self.detail_title = QLabel("항목을 선택하세요"); self.detail_title.setWordWrap(True); self.detail_title.setObjectName("sectionTitle")
        self.detail_meta = QLabel(""); self.detail_meta.setWordWrap(True)
        self.detail_history = QPlainTextEdit(); self.detail_history.setReadOnly(True); self.detail_history.setMinimumHeight(95); self.detail_history.setMaximumHeight(150)
        self.detail_vendor = QPlainTextEdit(); self.detail_vendor.setPlaceholderText("제보 페이지에서 확인한 벤더 / 프로그램 답변을 붙여넣으세요."); self.detail_vendor.setMinimumHeight(80); self.detail_vendor.setMaximumHeight(120)
        self.detail_notes = QPlainTextEdit(); self.detail_notes.setPlaceholderText("분석 메모, 재현 조건, 후속 계획 등"); self.detail_notes.setMinimumHeight(105)
        self.detail_vendor.setEnabled(False); self.detail_notes.setEnabled(False)
        self.detail_vendor.textChanged.connect(self._detail_edited)
        self.detail_notes.textChanged.connect(self._detail_edited)
        for widget in [self.detail_title, self.detail_meta, QLabel("변경 이력"), self.detail_history, QLabel("벤더 답변"), self.detail_vendor, QLabel("개인 메모"), self.detail_notes]: dl.addWidget(widget)
        self.detail_save_status = QLabel("메모와 벤더 답변은 입력 후 자동 저장됩니다.")
        dl.addWidget(self.detail_save_status)
        detail_scroll = QScrollArea(); detail_scroll.setObjectName("detailScroll")
        detail_scroll.setFrameShape(QFrame.NoFrame); detail_scroll.setWidgetResizable(True)
        detail_scroll.setWidget(detail_content); detail_layout.addWidget(detail_scroll, 1)
        br = QHBoxLayout()
        self.edit_button = QPushButton("정보 수정"); self.edit_button.clicked.connect(self.edit_selected); self.edit_button.setEnabled(False); br.addWidget(self.edit_button)
        self.save_button = QPushButton("지금 저장"); self.save_button.clicked.connect(self.save_detail); br.addWidget(self.save_button)
        self.open_button = QPushButton("제보 열기"); self.open_button.clicked.connect(self.open_selected_url); br.addWidget(self.open_button)
        self.delete_button = QPushButton("삭제"); self.delete_button.clicked.connect(self.delete_selected); br.addWidget(self.delete_button)
        for button in (self.save_button, self.open_button, self.delete_button):
            button.setEnabled(False)
        detail_layout.addLayout(br); splitter.addWidget(detail); splitter.setSizes([700,350]); l.addWidget(splitter)
        return w

    def _settings_page(self):
        w=QWidget(); l=QVBoxLayout(w); l.setContentsMargins(28,24,28,24)
        title=QLabel("GitHub 연결"); title.setObjectName("pageTitle"); l.addWidget(title)
        info=QLabel(
            "버그바운티 제보만 기록한다면 GitHub 연결은 필요하지 않습니다.\n"
            "GitHub Advisory를 동기화하려면 아래에서 로그인하세요. Client Secret은 필요하지 않습니다.\n"
            "미공개 Repository Security Advisory 조회를 위해 OAuth 'repo' 권한을 요청합니다. 이 권한은 저장소 접근 범위가 넓으므로 개인용 OAuth App만 사용하세요."
        ); info.setWordWrap(True); l.addWidget(info)

        oauth_box=QFrame(); oauth_box.setObjectName("card"); ob=QVBoxLayout(oauth_box)
        self.connection_label=QLabel("연결 상태 확인 중..."); self.connection_label.setObjectName("sectionTitle"); ob.addWidget(self.connection_label)
        form=QFormLayout()
        self.client_id_input=QLineEdit(); self.client_id_input.setPlaceholderText("예: Ov23li...")
        saved_client_id=self.settings.value("github/client_id", "", type=str); self.client_id_input.setText(saved_client_id)
        form.addRow("OAuth Client ID", self.client_id_input); ob.addLayout(form)
        auth_buttons=QHBoxLayout()
        login_btn=QPushButton("GitHub 로그인"); login_btn.clicked.connect(self.start_oauth_login)
        logout_btn=QPushButton("연결 해제"); logout_btn.clicked.connect(self.logout_github)
        register_btn=QPushButton("OAuth App 등록 페이지"); register_btn.clicked.connect(lambda: webbrowser.open("https://github.com/settings/applications/new"))
        auth_buttons.addWidget(login_btn); auth_buttons.addWidget(logout_btn); auth_buttons.addWidget(register_btn); auth_buttons.addStretch(1); ob.addLayout(auth_buttons)
        self.oauth_help=QLabel(
            "최초 1회: GitHub → Settings → Developer settings → OAuth Apps → New OAuth App에서 앱을 만들고 'Enable Device Flow'를 켠 뒤 Client ID만 입력하세요. "
            "Authorization callback URL은 Device Flow에서는 사용하지 않지만 등록 화면의 필수값이면 http://127.0.0.1 을 넣어도 됩니다."
        ); self.oauth_help.setWordWrap(True); ob.addWidget(self.oauth_help); l.addWidget(oauth_box)

        l.addWidget(QLabel("동기화할 저장소")); self.repo_list=QListWidget(); l.addWidget(self.repo_list,1)
        rr=QHBoxLayout(); self.repo_input=QLineEdit(); self.repo_input.setPlaceholderText("owner/repo 또는 https://github.com/owner/repo/security/advisories/GHSA-...")
        add=QPushButton("추가"); add.clicked.connect(self.add_repo); rem=QPushButton("선택 삭제"); rem.clicked.connect(self.remove_repo); rr.addWidget(self.repo_input,1); rr.addWidget(add); rr.addWidget(rem); l.addLayout(rr)
        self.sync_status=QLabel("준비"); l.addWidget(self.sync_status); test=QPushButton("저장소 동기화"); test.clicked.connect(self.sync_github); l.addWidget(test)
        self.refresh_connection_status()
        return w

    def _apply_style(self):
        self.setStyleSheet("""
            QWidget { background:#101317; color:#e8edf2; font-size:13px; }
            QLabel { background:transparent; }
            #nav { background:#171b21; border-right:1px solid #2a3038; }
            #nav QPushButton { text-align:left; padding:12px; border:none; border-radius:6px; }
            #nav QPushButton:hover, #nav QPushButton:checked { background:#252c35; }
            #brand { font-size:24px; font-weight:800; padding:16px 6px 24px 6px; }
            #pageTitle { font-size:24px; font-weight:700; margin-bottom:8px; }
            #sectionTitle { font-size:16px; font-weight:700; }
            #card { background:#171b21; border:1px solid #2a3038; border-radius:10px; }
            #statCard { background:#171b21; border:1px solid #2a3038; border-radius:10px; text-align:left; padding:0; }
            #statCard:hover, #statCard:focus { background:#222b35; border-color:#60a5fa; }
            #statCard QLabel { background:transparent; }
            #statValue { font-size:30px; font-weight:800; }
            #detail { background:#171b21; border:1px solid #2a3038; border-radius:8px; }
            #detailContent, #detailScroll { background:#171b21; }
            QPushButton { background:#242b33; border:1px solid #343d48; padding:8px 12px; border-radius:6px; }
            QPushButton:hover { background:#303945; }
            QPushButton:disabled { color:#78828d; background:#171b21; border-color:#252c35; }
            QLineEdit,QPlainTextEdit,QComboBox,QListWidget,QTableWidget { background:#0d1014; border:1px solid #303741; border-radius:5px; padding:6px; }
            QLineEdit:focus,QPlainTextEdit:focus,QComboBox:focus,QTableWidget:focus { border-color:#60a5fa; }
            QHeaderView::section { background:#171b21; color:#cfd7df; padding:7px; border:none; border-bottom:1px solid #303741; }
            QTableWidget { gridline-color:#20262d; }
            QTableWidget::item:alternate { background:#131820; }
            QStatusBar { background:#171b21; border-top:1px solid #2a3038; }
        """)

    def show_page(self, idx):
        self.stack.setCurrentIndex(idx)
        for i,b in enumerate([self.btn_dashboard,self.btn_reports,self.btn_settings]): b.setChecked(i==idx)
        if idx==2: self.refresh_repos()

    def _announce(self, message):
        self.statusBar().showMessage(message)
        if hasattr(self, "sync_status"):
            self.sync_status.setText(message)

    def _update_report_layout(self):
        compact = self.width() < 1100
        orientation = Qt.Vertical if compact else Qt.Horizontal
        if self.report_splitter.orientation() != orientation:
            self.report_splitter.setOrientation(orientation)
            self.report_splitter.setSizes([220, 320] if compact else [700, 350])

    def resizeEvent(self, event):
        super().resizeEvent(event)
        if hasattr(self, "report_splitter"):
            self._update_report_layout()

    def focus_search(self):
        self.show_page(1)
        self.search.setFocus()

    def clear_filters(self):
        self.show_reports_for()

    def show_reports_for(self, *, source="", state="", cve=False):
        if not self._save_detail_edits():
            return
        widgets = (self.search, self.source_filter, self.state_filter, self.cve_filter)
        for widget in widgets:
            widget.blockSignals(True)
        self.search.clear()
        self.source_filter.setCurrentIndex(self.source_filter.findData(source))
        states = BOUNTY_STATES if source == "bug_bounty" else GITHUB_STATES if source == "github" else ALL_STATES
        self.state_filter.clear()
        self.state_filter.addItems(["전체"] + states)
        self.state_filter.setCurrentText(state if state in states else "전체")
        self.cve_filter.setChecked(cve)
        for widget in widgets:
            widget.blockSignals(False)
        self.refresh_reports()
        self.show_page(1)

    def focus_selected_detail(self, _row, _column):
        if self.current_report_id:
            self.detail_notes.setFocus()

    def _source_filter_changed(self):
        previous = self.state_filter.currentText()
        source = self.source_filter.currentData()
        states = BOUNTY_STATES if source == "bug_bounty" else GITHUB_STATES if source == "github" else ALL_STATES
        self.state_filter.blockSignals(True)
        self.state_filter.clear()
        self.state_filter.addItems(["전체"] + states)
        self.state_filter.setCurrentText(previous if previous in states else "전체")
        self.state_filter.blockSignals(False)
        self.refresh_reports()

    def open_recent_report(self, row, _column):
        if not self._save_detail_edits():
            return
        report_id = int(self.recent_table.item(row, REPORT_ID_COLUMN).text())
        self.show_reports_for()
        self.current_report_id = report_id
        self.refresh_reports()

    def refresh_all(self):
        if not self._save_detail_edits(): return
        self.refresh_reports(); self.refresh_dashboard(); self.refresh_repos()

    def filtered_reports(self):
        rows=self.db.list_reports(); q=self.search.text().strip().lower() if hasattr(self,"search") else ""; st=self.state_filter.currentText() if hasattr(self,"state_filter") else "전체"; out=[]
        source = self.source_filter.currentData() if hasattr(self, "source_filter") else ""
        for r in rows:
            hay=" ".join(str(r[k] or "") for k in ["repository","title","ghsa_id","cve_id","state","severity","vendor_response","platform","program","report_reference","notes"]).lower()
            if q and q not in hay: continue
            if st!="전체" and r["state"]!=st: continue
            if source == "bug_bounty" and r["source"] != "bug_bounty": continue
            if source == "github" and r["source"] not in {"github", "manual"}: continue
            if hasattr(self, "cve_filter") and self.cve_filter.isChecked() and not r["cve_id"]: continue
            out.append(r)
        return out

    def change_label(self, r):
        events=self.db.list_events(r["id"],1)
        if not events: return "-"
        e=events[0]; m={"status":"STATUS","cve":"CVE","severity":"SEVERITY","github_update":"UPDATE","vendor_response":"VENDOR","created":"NEW"}
        return m.get(e["event_type"], e["event_type"].upper())

    def fill_table(self, table, rows):
        table.blockSignals(True)
        table.setRowCount(0)
        for r in rows:
            row=table.rowCount(); table.insertRow(row); change=self.change_label(r)
            vals=[TABLE_SOURCE_LABELS.get(r["source"], r["source"]),r["repository"],r["title"],r["state"],r["severity"] or "-",r["cve_id"] or "-",r["report_reference"] or "-",r["ghsa_id"] or "-",fmt_date(r["updated_at"]),change,str(r["id"])]
            for col,val in enumerate(vals):
                item=QTableWidgetItem(str(val)); item.setToolTip(SOURCE_LABELS.get(r["source"], r["source"]) if col == 0 else str(val)); table.setItem(row,col,item)
                if col==9 and change in {"STATUS","CVE","VENDOR"}: item.setFont(QFont("Segoe UI",10,QFont.Bold))
            if r["state"] == "published": table.item(row,3).setForeground(QBrush(QColor("#69db7c")))
            elif r["state"] == "accepted": table.item(row,3).setForeground(QBrush(QColor("#74c0fc")))
            elif r["state"] == "resolved": table.item(row,3).setForeground(QBrush(QColor("#63e6be")))
            elif r["state"] in {"closed", "duplicate", "not_applicable"}: table.item(row,3).setForeground(QBrush(QColor("#ff8787")))
            elif r["state"] in {"submitted", "triage","draft"}: table.item(row,3).setForeground(QBrush(QColor("#ffd43b")))
            if r["cve_id"]: table.item(row,5).setForeground(QBrush(QColor("#74c0fc")))
        table.blockSignals(False)

    def refresh_reports(self):
        if not hasattr(self, "report_table"): return
        selected_id = self.current_report_id
        rows = self.filtered_reports()
        total = len(self.db.list_reports())
        if total == 0:
            self.result_count.setText("제보가 없습니다. + 버그바운티 또는 + GitHub 제보로 시작하세요.")
        elif not rows:
            self.result_count.setText(f"검색 결과 0건 / 전체 {total}건 · CSV/Excel 내보내기 0건 · 필터를 확인하세요.")
        else:
            self.result_count.setText(f"검색 결과 {len(rows)}건 / 전체 {total}건 · CSV/Excel 내보내기 {len(rows)}건")
        self.fill_table(self.report_table, rows)
        for row in range(self.report_table.rowCount()):
            if int(self.report_table.item(row, REPORT_ID_COLUMN).text()) == selected_id:
                self.report_table.selectRow(row)
                return
        self.show_selected_report()

    def refresh_dashboard(self):
        rows=self.db.list_reports(); self.card_total.setText(str(len(rows))); self.card_triage.setText(str(sum(r["state"]=="triage" for r in rows))); self.card_draft.setText(str(sum(r["state"]=="draft" for r in rows))); self.card_published.setText(str(sum(r["state"]=="published" for r in rows))); self.card_cve.setText(str(sum(bool(r["cve_id"]) for r in rows)))
        self.card_bounty.setText(str(sum(r["source"] == "bug_bounty" for r in rows)))
        self.fill_table(self.recent_table,rows[:12])

    def add_report(self, checked=False, *, source="manual"):
        if not self._save_detail_edits(): return
        d=ReportDialog(self, source=source)
        while d.exec()==QDialog.Accepted:
            data=d.data()
            try:
                report_id = self.db.upsert_report(data)
            except (ValueError, sqlite3.Error) as e:
                QMessageBox.warning(self, "저장 실패", f"{e}\n입력한 내용은 남아 있습니다. 수정한 뒤 다시 저장하세요.")
                continue
            if data["source"] == "manual": self.db.add_repository(data["repository"])
            for widget in [self.search, self.source_filter, self.state_filter, self.cve_filter]: widget.blockSignals(True)
            self.search.clear(); self.source_filter.setCurrentIndex(0); self.state_filter.clear(); self.state_filter.addItems(["전체"] + ALL_STATES); self.cve_filter.setChecked(False)
            for widget in [self.search, self.source_filter, self.state_filter, self.cve_filter]: widget.blockSignals(False)
            self.current_report_id = report_id
            self.refresh_all(); self.show_page(1)
            self._announce("제보를 등록했습니다.")
            return

    def edit_selected(self):
        if not self._save_detail_edits(): return
        rid = self.current_report_id
        if not rid or rid != self.selected_report_id(): return
        report = self.db.get_report(rid)
        if not report: return
        if report["source"] == "github":
            QMessageBox.information(self, "GitHub 동기화 항목", "상태와 제보 정보는 GitHub에서 수정한 뒤 동기화하세요. 메모와 벤더 답변은 이 화면에서 저장할 수 있습니다.")
            return
        dialog = ReportDialog(self, report=report)
        while dialog.exec() == QDialog.Accepted:
            data = dialog.data()
            try:
                self.db.update_report(rid, data)
            except (ValueError, sqlite3.Error) as e:
                QMessageBox.warning(self, "수정 실패", f"{e}\n입력한 내용은 남아 있습니다. 수정한 뒤 다시 저장하세요.")
                continue
            if data["source"] == "manual": self.db.add_repository(data["repository"])
            self.refresh_all(); self._announce("제보 정보를 수정했습니다.")
            return

    def add_repo_from_link_dialog(self):
        self.show_page(2); self.repo_input.setFocus()

    def selected_report_id(self):
        if not self.report_table.selectedItems(): return None
        return int(self.report_table.item(self.report_table.currentRow(),REPORT_ID_COLUMN).text())

    def show_selected_report(self):
        rid=self.selected_report_id()
        if self._detail_dirty and not self._save_detail_edits():
            self._restore_selection_after_save_error()
            return
        self.current_report_id=rid
        self.edit_button.setEnabled(False)
        self.save_button.setEnabled(bool(rid))
        self.delete_button.setEnabled(bool(rid))
        self.detail_notes.setEnabled(bool(rid))
        self.detail_vendor.setEnabled(bool(rid))
        self._loading_detail = True
        if not rid:
            self.detail_title.setText("항목을 선택하세요"); self.detail_meta.clear()
            self.detail_notes.clear(); self.detail_vendor.clear(); self.detail_history.clear()
            self.open_button.setEnabled(False)
            self.detail_save_status.setText("항목을 선택하면 메모를 입력할 수 있습니다.")
            self._loading_detail = False
            return
        r=self.db.get_report(rid); self.detail_title.setText(r["title"])
        self.edit_button.setEnabled(r["source"] != "github")
        self.open_button.setEnabled(bool(r["html_url"]))
        meta = [f"상태: {r['state']}  |  Severity: {r['severity'] or '-'}", r["repository"],
                f"출처: {SOURCE_LABELS.get(r['source'], r['source'])}"]
        if r["source"] == "bug_bounty":
            meta += [f"플랫폼: {r['platform'] or '-'}", f"프로그램: {r['program'] or '-'}", f"제보 번호: {r['report_reference'] or '-'}"]
        else:
            meta.append(f"GHSA: {r['ghsa_id'] or '-'}")
        meta += [f"CVE: {r['cve_id'] or '-'}", f"등록: {fmt_date(r['submitted_at'])}", f"수정: {fmt_date(r['updated_at'])}"]
        if r["synced_at"]: meta.append(f"GitHub 동기화: {fmt_date(r['synced_at'])}")
        self.detail_meta.setText("\n".join(meta))
        self.detail_notes.setPlainText(r["notes"] or ""); self.detail_vendor.setPlainText(r["vendor_response"] or "")
        self._refresh_detail_history(rid)
        self.detail_save_status.setText("메모와 벤더 답변은 입력 후 자동 저장됩니다.")
        self._loading_detail = False

    def _restore_selection_after_save_error(self):
        previous_row = None
        for row in range(self.report_table.rowCount()):
            item = self.report_table.item(row, REPORT_ID_COLUMN)
            if item and int(item.text()) == self.current_report_id:
                previous_row = row
                break
        self.report_table.blockSignals(True)
        if previous_row is None:
            self.report_table.clearSelection()
            for button in (self.edit_button, self.open_button, self.delete_button):
                button.setEnabled(False)
        else:
            self.report_table.selectRow(previous_row)
        self.report_table.blockSignals(False)
        self._announce("메모 저장에 실패해 제보 선택을 바꾸지 않았습니다. 다시 저장하세요.")

    def _refresh_detail_history(self, report_id):
        lines=[]
        for e in self.db.list_events(report_id,30):
            old=(e["old_value"] or "-").replace("\n"," ")[:80]; new=(e["new_value"] or "-").replace("\n"," ")[:80]
            lines.append(f"{fmt_date(e['created_at'])}  {e['event_type']}: {old} → {new}")
        self.detail_history.setPlainText("\n".join(lines) if lines else "변경 이력 없음")

    def _detail_edited(self):
        if self._loading_detail or not self.current_report_id:
            return
        self._detail_dirty = True
        self.detail_save_status.setText("자동 저장 대기 중...")
        self._detail_save_timer.start()

    def _save_detail_edits(self):
        if not self._detail_dirty or not self.current_report_id:
            return True
        self._detail_save_timer.stop()
        try:
            self.db.update_notes(self.current_report_id, self.detail_notes.toPlainText())
            self.db.update_vendor_response(self.current_report_id, self.detail_vendor.toPlainText())
        except (ValueError, sqlite3.Error) as error:
            self.detail_save_status.setText("저장 실패 · 다시 시도하세요.")
            QMessageBox.warning(self, "메모 저장 실패", str(error))
            return False
        self._detail_dirty = False
        self.detail_save_status.setText("저장됨")
        return True

    def _autosave_detail(self):
        if self._save_detail_edits() and self.current_report_id:
            self._refresh_detail_history(self.current_report_id)
            self.refresh_dashboard()

    def save_detail(self):
        if not self.current_report_id: return
        if self._save_detail_edits():
            self.refresh_all()
            self._announce("메모와 벤더 답변이 저장됐습니다.")

    def open_selected_url(self):
        rid=self.current_report_id
        if not rid or rid != self.selected_report_id(): return
        r=self.db.get_report(rid)
        if r["html_url"]: webbrowser.open(r["html_url"])
        else: QMessageBox.information(self,"URL 없음","등록된 제보 URL이 없습니다.")

    def delete_selected(self):
        rid=self.current_report_id
        if not rid or rid != self.selected_report_id():
            return
        report = self.db.get_report(rid)
        if not report:
            return
        message = f"'{report['title']}'의 로컬 기록과 변경 이력을 삭제할까요?"
        if report["source"] == "github":
            message += "\n\nGitHub 동기화 항목은 다음 동기화 때 다시 등록될 수 있습니다. 계속 제외하려면 GitHub 설정에서 저장소도 제거하세요."
        elif report["source"] == "manual":
            message += "\n\n등록된 GitHub 저장소는 동기화 목록에 남습니다."
        if QMessageBox.question(self,"제보 삭제",message)==QMessageBox.Yes:
            self._detail_save_timer.stop()
            self._detail_dirty = False
            self.db.delete_report(rid); self.current_report_id=None; self.refresh_all()
            self._announce("제보를 삭제했습니다.")

    def refresh_connection_status(self):
        try:
            token=keyring.get_password(KEYRING_SERVICE,KEYRING_USER)
        except Exception:
            token=None
        if not hasattr(self, "connection_label"):
            return
        if not token:
            self.connection_label.setText("● GitHub 연결 안 됨")
            return
        self.connection_label.setText("● GitHub 로그인 정보 저장됨 · 동기화 시 연결을 확인합니다")

    def start_oauth_login(self):
        client_id=self.client_id_input.text().strip()
        if not client_id:
            QMessageBox.warning(self,"Client ID","GitHub OAuth App의 Client ID를 입력하세요.")
            return
        self.settings.setValue("github/client_id", client_id)
        if self.auth_worker and self.auth_worker.isRunning():
            self._announce("GitHub 로그인이 이미 진행 중입니다.")
            return
        self.sync_status.setText("GitHub 로그인 코드 요청 중...")
        self.auth_worker=OAuthWorker(client_id)
        self.auth_worker.code_ready.connect(self._oauth_code_ready)
        self.auth_worker.authenticated.connect(self._oauth_authenticated)
        self.auth_worker.failed.connect(self._oauth_failed)
        self.auth_worker.start()

    def _oauth_code_ready(self, verification_uri, user_code, expires_in):
        if self._closing_pending: return
        QApplication.clipboard().setText(user_code)
        webbrowser.open(verification_uri)
        self.sync_status.setText(f"브라우저에서 인증하세요. 코드 {user_code} (클립보드 복사됨)")
        QMessageBox.information(
            self,"GitHub 로그인",
            f"브라우저에서 GitHub 인증 페이지를 열었습니다.\n\n인증 코드: {user_code}\n\n코드는 클립보드에 복사했습니다. {expires_in//60}분 안에 승인하세요."
        )

    def _oauth_authenticated(self, login, token, refresh_token):
        if self._closing_pending: return
        try:
            keyring.set_password(KEYRING_SERVICE,KEYRING_USER,token)
            if refresh_token:
                keyring.set_password(KEYRING_SERVICE,KEYRING_REFRESH,refresh_token)
        except Exception as e:
            QMessageBox.critical(self,"저장 실패",f"OAuth 토큰을 Windows 자격 증명 저장소에 저장하지 못했습니다.\n{e}")
            return
        self.connection_label.setText(f"● GitHub 연결됨: @{login}")
        self.sync_status.setText(f"GitHub 로그인 완료: @{login}")
        QMessageBox.information(self,"GitHub 연결 완료",f"@{login} 계정과 연결했습니다.\n이제 저장소를 추가하고 동기화하면 됩니다.")

    def _oauth_failed(self, message):
        if self._closing_pending: return
        self.sync_status.setText("GitHub 로그인 실패")
        QMessageBox.critical(self,"GitHub OAuth 오류",message)

    def logout_github(self):
        try:
            keyring.delete_password(KEYRING_SERVICE,KEYRING_USER)
        except Exception:
            pass
        try:
            keyring.delete_password(KEYRING_SERVICE,KEYRING_REFRESH)
        except Exception:
            pass
        self.connection_label.setText("● GitHub 연결 안 됨")
        self.sync_status.setText("GitHub 연결 해제됨")

    def add_repo(self):
        raw=self.repo_input.text().strip(); repo=repository_from_input(raw)
        if not repo: QMessageBox.warning(self,"형식","owner/repository 또는 GitHub repository/advisory URL을 입력하세요."); return
        self.db.add_repository(repo); self.repo_input.clear(); self.refresh_repos(); self.sync_status.setText(f"저장소 추가: {repo}")

    def remove_repo(self):
        item=self.repo_list.currentItem()
        if item: self.db.remove_repository(item.text()); self.refresh_repos()

    def refresh_repos(self):
        if hasattr(self,"repo_list"): self.repo_list.clear(); self.repo_list.addItems(self.db.list_repositories())

    def sync_github(self):
        if not self._save_detail_edits(): return
        try: token=keyring.get_password(KEYRING_SERVICE,KEYRING_USER)
        except Exception: token=None
        repos=self.db.list_repositories()
        if not token: self.show_page(2); QMessageBox.information(self,"GitHub 연결","먼저 GitHub 로그인 버튼으로 계정을 연결하세요."); return
        if not repos: self.show_page(2); QMessageBox.information(self,"저장소","동기화할 저장소를 하나 이상 추가하세요."); return
        if self.worker and self.worker.isRunning():
            self._announce("GitHub 동기화가 이미 진행 중입니다.")
            return
        self._announce("GitHub 연결 중..."); self.worker=SyncWorker(token,repos); self.worker.progress.connect(self._announce); self.worker.item.connect(self._receive_sync_item); self.worker.done.connect(self._sync_done); self.worker.failed.connect(self._sync_failed); self.worker.start()

    def _receive_sync_item(self,data): self.db.upsert_report(data)
    def _sync_done(self,total,failures):
        if self._closing_pending: return
        self.refresh_all(); self._announce(f"동기화 완료: {total}건 반영 / 실패 {failures}개"); QMessageBox.information(self,"동기화 완료",f"내 advisory {total}건을 반영했습니다.\n저장소 오류: {failures}개")
    def _sync_failed(self,message):
        if self._closing_pending: return
        self._announce("동기화 실패"); QMessageBox.critical(self,"GitHub 오류",message)

    def _export_rows(self):
        if self.stack.currentIndex() != 1:
            return self.db.list_reports()
        return self.filtered_reports()

    def export_csv(self):
        if not self._save_detail_edits(): return
        path,_=QFileDialog.getSaveFileName(self,"CSV 내보내기","vulnerability_reports.csv","CSV (*.csv)")
        if not path: return
        rows = self._export_rows()
        fields=["repository","title","state","severity","cve_id","ghsa_id","submitted_at","updated_at","last_change_at","vendor_response_at","vendor_response","notes","html_url","source","platform","program","report_reference"]
        with open(path,"w",newline="",encoding="utf-8-sig") as f:
            w=csv.writer(f); w.writerow(fields)
            for r in rows: w.writerow([r[k] if k in r.keys() else "" for k in fields])
        self._announce(f"CSV {len(rows)}건 저장: {path}")

    def export_xlsx(self):
        if not self._save_detail_edits(): return
        path,_=QFileDialog.getSaveFileName(self,"Excel 내보내기","vulnerability_reports.xlsx","Excel (*.xlsx)")
        if not path: return
        wb=Workbook(); ws=wb.active; ws.title="Vulnerability Reports"
        fields=["Target / Repository","Title","Status","Severity","CVE","GHSA","Submitted","Updated","Last Change","Vendor Response At","Vendor Response","Notes","URL","Source","Platform","Program","Report Reference"]
        ws.append(fields)
        for c in ws[1]: c.font=Font(bold=True)
        keys=["repository","title","state","severity","cve_id","ghsa_id","submitted_at","updated_at","last_change_at","vendor_response_at","vendor_response","notes","html_url","source","platform","program","report_reference"]
        rows = self._export_rows()
        for r in rows: ws.append([r[k] if k in r.keys() else "" for k in keys])
        ws.freeze_panes="A2"; ws.auto_filter.ref=ws.dimensions
        widths=[28,52,18,12,20,24,22,22,22,22,45,45,55,16,20,28,24]
        for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=w
        ev=wb.create_sheet("Change History"); ev.append(["Report ID","Target / Repository","Source","Platform","Program","Report Reference","GHSA","Event","Old","New","Time"])
        for c in ev[1]: c.font=Font(bold=True)
        for r in rows:
            for e in reversed(self.db.list_events(r["id"],1000)):
                ev.append([r["id"],r["repository"],r["source"],r["platform"],r["program"],r["report_reference"],r["ghsa_id"],e["event_type"],e["old_value"],e["new_value"],e["created_at"]])
        ev.freeze_panes="A2"; wb.save(path); self._announce(f"Excel {len(rows)}건 저장: {path}")

    def closeEvent(self,event):
        if not self._save_detail_edits():
            event.ignore()
            return
        running = [thread for thread in (self.auth_worker, self.worker)
                   if thread and thread.isRunning()]
        if running:
            if self.auth_worker and self.auth_worker.isRunning():
                self.auth_worker.cancel()
            if not self._closing_pending:
                self._closing_pending = True
                for thread in running:
                    thread.finished.connect(self._finish_pending_close)
            self._announce("진행 중인 GitHub 작업이 끝나면 종료합니다...")
            event.ignore()
            return
        self.db.close(); super().closeEvent(event)

    def _finish_pending_close(self):
        if not any(thread and thread.isRunning() for thread in (self.auth_worker, self.worker)):
            QTimer.singleShot(0, self.close)


def main():
    app=QApplication(sys.argv); app.setApplicationName("Vulnerability Tracker"); app.setFont(QFont("Segoe UI",10)); win=MainWindow(); win.show(); sys.exit(app.exec())


if __name__=="__main__": main()
